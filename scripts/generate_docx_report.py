import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def create_document():
    doc = Document()

    # Page Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Styles
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # Primary Colors
    c_primary = RGBColor(0x1B, 0x36, 0x5D)    # Deep Navy
    c_secondary = RGBColor(0x2B, 0x6C, 0xB0)  # Steel Blue
    c_dark = RGBColor(0x22, 0x22, 0x22)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("Especificação Técnica da Rede Neural e Treinamento")
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = c_primary

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(20)
    run_sub = p_sub.add_run("Projeto InterpretArte – Tradutor de Libras em Tempo Real para Área da Saúde")
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(13)
    run_sub.font.italic = True
    run_sub.font.color.rgb = c_secondary

    # Divider Box / Metadata
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_data = [
        [("Documento:", True), ("Especificação de Arquitetura e Treinamento de IA", False)],
        [("Framework:", True), ("TensorFlow 2.x / Keras com MediaPipe Holistic", False)],
    ]
    for row_idx, row in enumerate(meta_table.rows):
        for col_idx, cell in enumerate(row.cells):
            set_cell_background(cell, "F0F4F8")
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            cell.width = Inches(3.25)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            prefix, bold = meta_data[row_idx][col_idx]
            r = p.add_run(prefix + " ")
            r.font.bold = True
            r.font.size = Pt(10)
            r.font.color.rgb = c_primary
            if not bold:
                pass

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    def add_heading_1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(15)
        r.font.bold = True
        r.font.color.rgb = c_primary
        return h

    def add_heading_2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(12.5)
        r.font.bold = True
        r.font.color.rgb = c_secondary
        return h

    # Section 1
    add_heading_1("1. Visão Geral e Estrutura dos Arquivos de Código")
    p = doc.add_paragraph(
        "O treinamento e a inferência dos sinais em Língua Brasileira de Sinais (Libras) "
        "são implementados em Python utilizando o framework TensorFlow / Keras em conjunto com a biblioteca "
        "Google MediaPipe Holistic para extração das coordenadas dos membros corporais em vídeos. "
        "O fluxo de engenharia de dados e modelagem está modularizado nos seguintes arquivos no repositório:"
    )
    p.paragraph_format.space_after = Pt(6)

    bullets = [
        ("src/feature_extraction.py: ", "Responsável por processar os vídeos gravados e extrair as coordenadas tridimensionais (keypoints) de pose e mãos via MediaPipe Holistic, salvando os arrays numpy (.npy)."),
        ("src/preprocess_data.py: ", "Padroniza o comprimento das sequências temporais (fixado em 45 frames por corte) utilizando padding e gera as divisões de treino, validação e teste (X_train, y_train, etc.)."),
        ("src/model_training.py: ", "Contém a definição exata da rede neural profunda recorrente (LSTM), compilação com otimizador Adam, callbacks de regularização, pesos de classes e rotina completa de treino."),
        ("scripts/evaluate_model.py: ", "Executa a avaliação estatística no conjunto de teste independente, gerando matriz de confusão e métricas de acurácia, precisão, revocação e F1-Score."),
        ("scripts/run_full_pipeline.py: ", "Orquestrador de execução em lote do pipeline de ponta a ponta (corte -> extração -> pré-processamento -> treino -> avaliação).")
    ]
    for b_title, b_desc in bullets:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(3)
        r_bt = bp.add_run(b_title)
        r_bt.font.bold = True
        r_bt.font.color.rgb = c_dark
        r_bd = bp.add_run(b_desc)

    # Section 2
    add_heading_1("2. Formato e Dimensões dos Dados de Entrada (Input Shape)")
    p2 = doc.add_paragraph(
        "A rede neural recebe sequências temporais tridimensionais no formato (Batch_Size, Frames, Features), "
        "onde cada amostra representa a execução gestual contínua de um sinal:"
    )
    p2.paragraph_format.space_after = Pt(4)

    input_details = [
        ("Formato de Entrada (Input Shape): ", "(None, 45, 258)"),
        ("Número de Frames por Sequência: ", "45 frames normalizados (com padding para sequências menores)."),
        ("Dimensão de Características por Frame: ", "258 valores contínuos normalizados:")
    ]
    for k, v in input_details:
        p_item = doc.add_paragraph(style='List Bullet')
        p_item.paragraph_format.space_after = Pt(2)
        r_k = p_item.add_run(k)
        r_k.font.bold = True
        r_v = p_item.add_run(v)

    # Sub-bullets for 258 features
    features_desc = [
        ("Pose Corporal: ", "33 marcos corporais × 4 valores (x, y, z, visibilidade) = 132 dimensões."),
        ("Mão Esquerda: ", "21 marcos articulares × 3 coordenadas espaciais (x, y, z) = 63 dimensões."),
        ("Mão Direita: ", "21 marcos articulares × 3 coordenadas espaciais (x, y, z) = 63 dimensões."),
        ("Total consolidado por frame: ", "132 + 63 + 63 = 258 dimensões numéricas por passo de tempo.")
    ]
    for fk, fv in features_desc:
        pf = doc.add_paragraph(style='List Bullet 2')
        pf.paragraph_format.space_after = Pt(2)
        r_fk = pf.add_run(fk)
        r_fk.font.bold = True
        r_fv = pf.add_run(fv)

    # Section 3
    add_heading_1("3. Arquitetura da Rede Neural (Sequential LSTM + Dense)")
    p3 = doc.add_paragraph(
        "A arquitetura foi projetada com uma combinação de camadas recorrentes Long Short-Term Memory (LSTM), "
        "adequadas para capturar a dinâmica espaço-temporal dos sinais em Libras, seguidas por camadas densas "
        "(Fully Connected) de tomada de decisão e uma camada Softmax final de classificação probabilística."
    )
    p3.paragraph_format.space_after = Pt(6)

    # Table of Layers
    table = doc.add_table(rows=7, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    headers = ["#", "Camada", "Tipo", "Unidades", "Ativação", "Dimensão de Saída", "Parâmetros"]
    # Adjust table col count
    table._tbl.remove(table.rows[0]._tr) # recreate properly

    table = doc.add_table(rows=1, cols=7)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    col_widths = [Inches(0.4), Inches(1.1), Inches(1.0), Inches(0.9), Inches(0.9), Inches(1.5), Inches(1.1)]

    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "1B365D")
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=80, right=80)
        hdr_cells[i].paragraphs[0].runs[0].font.bold = True
        hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        hdr_cells[i].paragraphs[0].runs[0].font.size = Pt(9.5)
        hdr_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    layers_data = [
        ("1", "lstm", "LSTM", "64", "ReLU", "(None, 45, 64)", "82.688"),
        ("2", "lstm_1", "LSTM", "128", "ReLU", "(None, 45, 128)", "98.816"),
        ("3", "lstm_2", "LSTM", "64", "ReLU", "(None, 64)", "49.408"),
        ("4", "dense", "Dense", "64", "ReLU", "(None, 64)", "4.160"),
        ("5", "dense_1", "Dense", "32", "ReLU", "(None, 32)", "2.080"),
        ("6", "dense_2", "Dense", "5", "Softmax", "(None, 5)", "165"),
    ]

    for row_idx, data in enumerate(layers_data):
        row_cells = table.add_row().cells
        bg_color = "F9FBFC" if row_idx % 2 == 0 else "FFFFFF"
        for c_idx, val in enumerate(data):
            row_cells[c_idx].text = val
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=80, right=80)
            p_cell = row_cells[c_idx].paragraphs[0]
            p_cell.paragraph_format.space_after = Pt(0)
            if len(p_cell.runs) > 0:
                p_cell.runs[0].font.size = Pt(9.5)
                p_cell.runs[0].font.color.rgb = c_dark
                if c_idx in [0, 2, 3, 4]:
                    p_cell.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif c_idx == 6:
                    p_cell.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    # Fix widths
    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    p_params = doc.add_paragraph()
    p_params.paragraph_format.space_before = Pt(6)
    p_params.paragraph_format.space_after = Pt(10)
    r_pt = p_params.add_run("Total de Parâmetros Treináveis: ")
    r_pt.font.bold = True
    p_params.add_run("237.317 parâmetros (aprox. 927.02 KB em memória). Não há parâmetros congelados (0 non-trainable).")

    # Section 4
    add_heading_1("4. Hiperparâmetros de Treinamento e Otimizador")

    hyper_params = [
        ("Otimizador: ", "Adam (Adaptive Moment Estimation)."),
        ("Taxa de Aprendizado (Learning Rate): ", "0.0003 (3e-4) inicial, proporcionando convergência suave sem oscilações drásticas."),
        ("Clipagem de Gradiente (clipnorm): ", "1.0 — técnica essencial para mitigar o problema de explosão de gradientes (exploding gradients) em redes recorrentes LSTM."),
        ("Função de Perda (Loss): ", "Categorical Crossentropy (adequada para problemas multiclasse com saídas one-hot encoded)."),
        ("Métrica Principal: ", "Categorical Accuracy (percentual de predições corretas da classe alvo)."),
        ("Tamanho do Lote (Batch Size): ", "32 amostras por iteração."),
        ("Épocas de Treinamento (Epochs): ", "Até 200 épocas (com critério de parada antecipada)."),
        ("Ponderação de Classes (class_weight): ", "Cálculo balanceado dinamicamente com base na frequência inversa das amostras no split de treino, acrescido de reforço multiplicador de 1.5× para as classes de maior confusão gestual (curar e doenca).")
    ]
    for hk, hv in hyper_params:
        hp = doc.add_paragraph(style='List Bullet')
        hp.paragraph_format.space_after = Pt(3)
        r_hk = hp.add_run(hk)
        r_hk.font.bold = True
        r_hk.font.color.rgb = c_dark
        hp.add_run(hv)

    # Section 5
    add_heading_1("5. Callbacks e Mecanismos de Regularização")
    p5 = doc.add_paragraph(
        "Durante o ciclo de ajuste dos pesos, três callbacks do Keras atuam de forma automática para garantir a generalização do modelo e evitar overfitting:"
    )
    p5.paragraph_format.space_after = Pt(4)

    callbacks_info = [
        ("ModelCheckpoint: ", "Monitora 'val_categorical_accuracy' com modo 'max' e 'save_best_only=True'. Garante que o arquivo models/melhor_modelo.keras retenha sempre a melhor versão da rede na validação, prevenindo regressões de performance."),
        ("EarlyStopping: ", "Monitora 'val_loss' com paciência de 15 épocas ('patience=15') e restaura os melhores pesos automaticamente ('restore_best_weights=True'), interrompendo o treino quando a rede atinge platô."),
        ("ReduceLROnPlateau: ", "Monitora 'val_loss'. Se a perda estagnar por 8 épocas ('patience=8'), reduz a taxa de aprendizado pela metade ('factor=0.5') até o piso de '1e-5', permitindo ajustes finos nos pesos.")
    ]
    for ck, cv in callbacks_info:
        cp = doc.add_paragraph(style='List Bullet')
        cp.paragraph_format.space_after = Pt(3)
        r_ck = cp.add_run(ck)
        r_ck.font.bold = True
        r_ck.font.color.rgb = c_dark
        cp.add_run(cv)

    # Section 6
    add_heading_1("6. Conjunto de Dados, Classes e Divisão Experimental")
    p6 = doc.add_paragraph(
        "O vocabulário atual compreende 5 sinais prioritários voltados para o atendimento clínico e hospitalar em Libras:"
    )
    p6.paragraph_format.space_after = Pt(4)

    classes_list = [
        ("1. curar", "Sinal indicativo de cura / recuperação clínica."),
        ("2. doenca", "Sinal indicativo de patologia / enfermidade."),
        ("3. dor", "Sinal indicativo de queixa álgica / sensação de dor do paciente."),
        ("4. remedio", "Sinal indicativo de prescrição / fármaco / medicação."),
        ("5. saude", "Sinal indicativo de estado geral de saúde / bem-estar.")
    ]
    for c_name, c_desc in classes_list:
        cp = doc.add_paragraph(style='List Bullet')
        cp.paragraph_format.space_after = Pt(2)
        r_cn = cp.add_run(c_name + ": ")
        r_cn.font.bold = True
        cp.add_run(c_desc)

    p_splits = doc.add_paragraph()
    p_splits.paragraph_format.space_before = Pt(6)
    p_splits.paragraph_format.space_after = Pt(4)
    p_splits.add_run("Divisão das sequências de treino, validação e teste:").font.bold = True

    splits_data = [
        ("Conjunto de Treinamento (X_train / y_train): ", "213 sequências de vídeo (70.5%)"),
        ("Conjunto de Validação (X_val / y_val): ", "41 sequências de vídeo (13.6%)"),
        ("Conjunto de Teste Independente (X_test / y_test): ", "48 sequências de vídeo (15.9%)"),
        ("Total de Amostras Avaliadas: ", "302 sequências de 45 frames cada.")
    ]
    for sk, sv in splits_data:
        sp = doc.add_paragraph(style='List Bullet')
        sp.paragraph_format.space_after = Pt(2)
        r_sk = sp.add_run(sk)
        r_sk.font.bold = True
        sp.add_run(sv)

    # Save
    out_path = os.path.abspath("DOCUMENTO_TECNICO_REDE_NEURAL_LIBRAS.docx")
    doc.save(out_path)
    print(f"Documento salvo com sucesso em: {out_path}")

if __name__ == "__main__":
    create_document()
